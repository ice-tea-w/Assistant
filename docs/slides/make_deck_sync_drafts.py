# -*- coding: utf-8 -*-
"""设备同步功能 —— 三个视觉方向初稿（每方向 2 页，用于方向选择）

用法： python make_deck_sync_drafts.py
产出： docs/slides/_drafts/draft_a_mirror.pptx
       docs/slides/_drafts/draft_b_duo.pptx
       docs/slides/_drafts/draft_c_flow.pptx

三个方向都建立在既有配色（NAVY #142238 / BLUE #4A90D9）之上，差异在视觉母题：
  A 镜像 —— 延续前三套 deck 的模板语言，用左右对称表达"两台设备"
  B 双栏 —— 把"两个对端 + 中间通道"做成版式骨架，突出分布式结构
  C 状态流 —— 以文件生命周期时间线为骨架，把"冲突可见"做成主角

两条硬约束（由 check_deck_geometry.py / check_deck_contrast.py 验证）：
  1. 文本框宽度一律用 tw() 估算后设定，不"猜"。框宽过大有两个后果：
     文字压到不该压的区域；对比度自检因中心点落在错误背景上而误报。
  2. 文字颜色用 *_TXT 深色变体而非主色。主色（BLUE/TEAL/AMBER/GREEN）
     是为色块和图形设计的，直接当小字用达不到 4.5:1 对比度。
"""
import os

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

FONT    = "Microsoft YaHei"

# ---- 图形/色块用色（沿用既有模板）----
NAVY    = RGBColor(0x14, 0x22, 0x38)
NAVY2   = RGBColor(0x1E, 0x30, 0x4C)
NAVY3   = RGBColor(0x2A, 0x3F, 0x5E)
NAVY4   = RGBColor(0x1B, 0x3A, 0x4C)
BLUE    = RGBColor(0x4A, 0x90, 0xD9)
BLUE_L  = RGBColor(0xE8, 0xF1, 0xFB)
TEAL    = RGBColor(0x1F, 0x7A, 0x8C)
TEAL_L  = RGBColor(0xE4, 0xF2, 0xF4)
GREEN   = RGBColor(0x2E, 0x8B, 0x57)
GREEN_L = RGBColor(0xE7, 0xF4, 0xEC)
AMBER   = RGBColor(0xB5, 0x6E, 0x00)
AMBER_L = RGBColor(0xFD, 0xF3, 0xE2)
RED     = RGBColor(0xC0, 0x39, 0x2B)
RED_L   = RGBColor(0xFB, 0xEC, 0xEA)

INK     = RGBColor(0x1A, 0x1A, 0x1A)
MUTED   = RGBColor(0x5E, 0x6A, 0x7A)
MUTED_L = RGBColor(0x8E, 0x9A, 0xAD)
LINE    = RGBColor(0xD8, 0xDE, 0xE6)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
BG      = RGBColor(0xF7, 0xF9, 0xFC)
ZEBRA   = RGBColor(0xF1, 0xF5, 0xFA)
SOFT    = RGBColor(0xB9, 0xC6, 0xD8)   # 深色底上的正文灰
SOFT2   = RGBColor(0xC7, 0xD4, 0xE6)   # 深色底上的次级灰

# ---- 文字专用深色变体（保证 >=4.5:1）----
BLUE_TXT  = RGBColor(0x2C, 0x5F, 0x9E)
TEAL_TXT  = RGBColor(0x17, 0x62, 0x70)
GREEN_TXT = RGBColor(0x26, 0x72, 0x4A)
AMBER_TXT = RGBColor(0x8A, 0x54, 0x00)

# ---- 色块填充的加深变体：白字压主色只有 4.0~4.3:1，不够 4.5 门槛。
#      加深明度、保持色相，语义（绿=正常 / 琥珀=注意）不变。
GREEN_CHIP = RGBColor(0x2A, 0x7D, 0x50)
AMBER_CHIP = RGBColor(0x9C, 0x5D, 0x00)

W, H = 13.333, 7.5
M, CW = 0.75, 11.833
BODY_TOP = 1.80

OUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_drafts")


# ---------------------------------------------------------------- 尺寸估算

def tw(text, size, bold=False):
    """估算一段文字单行渲染所需宽度（英寸）。

    中文按 1.0 em、英文数字按 0.55 em 估。用于设定文本框宽度，
    避免"框远大于文字"导致的排版与自检问题。
    """
    best = 0.0
    for line in str(text).split("\n"):
        cjk = sum(1 for ch in line if ord(ch) > 0x2E80)
        other = len(line) - cjk
        w = (cjk * size + other * size * 0.55) * (1.03 if bold else 1.0)
        best = max(best, w)
    return best / 72.0 + 0.08


# ---------------------------------------------------------------- 版式基础

def style_font(run, size=None, bold=None, color=None, name=FONT, italic=None):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ('a:ea', 'a:cs'):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set('typeface', name)
    if size   is not None: run.font.size  = Pt(size)
    if bold   is not None: run.font.bold  = bold
    if italic is not None: run.font.italic = italic
    if color  is not None: run.font.color.rgb = color


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def tbox(slide, x, y, text, size, bold=False, color=INK, anchor=MSO_ANCHOR.TOP,
         align=PP_ALIGN.LEFT, h=None, maxw=None):
    """宽度按文字实际长度设定的文本框（首选）。"""
    w = tw(text, size, bold)
    if maxw is not None:
        w = min(w, maxw)
    if h is None:
        h = size * 1.45 / 72.0
    tf = textbox(slide, x, y, w, h, anchor=anchor)
    para(tf, text, size=size, bold=bold, color=color, first=True,
         align=align, after=0)
    return tf


def para(tf, text, size=14, bold=False, color=INK, first=False, before=0, after=5,
         align=PP_ALIGN.LEFT, line=1.2, italic=False, font=FONT):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(before); p.space_after = Pt(after); p.line_spacing = line
    r = p.add_run(); r.text = text
    style_font(r, size=size, bold=bold, color=color, italic=italic, name=font)
    return p


def rect(slide, x, y, w, h, fill=None, line_color=None, line_w=0.75, radius=None):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None: shp.fill.background()
    else:            shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line_color is None: shp.line.fill.background()
    else:                  shp.line.color.rgb = line_color; shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    if radius: shp.adjustments[0] = radius
    return shp


def chip(slide, x, y, text, fill, txt_color, size=12, pad=0.30, h=0.36, radius=0.4):
    """按文字长度自适应的胶囊标签。返回 (x, y, w, h)。"""
    w = tw(text, size, True) + pad * 2
    rect(slide, x, y, w, h, fill=fill, radius=radius)
    para(textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.MIDDLE), text, size=size,
         bold=True, color=txt_color, first=True, align=PP_ALIGN.CENTER, after=0)
    return x, y, w, h


def blank(prs, bg=BG):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, fill=bg)
    return s


def header(slide, kicker, title, dark=False):
    tbox(slide, M, 0.42, kicker, 12, bold=True, color=BLUE_TXT if not dark else BLUE)
    tbox(slide, M, 0.76, title, 27, bold=True, color=WHITE if dark else INK)
    rect(slide, M, 1.50, 1.05, 0.045, fill=BLUE)


def pageno(slide, n, dark=False):
    t = f"{n:02d}"
    tbox(slide, W - M - tw(t, 11, True), 6.98, t, 11, bold=True,
         color=MUTED_L if dark else MUTED)


def deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    return prs


def save(prs, name):
    os.makedirs(OUT_DIR, exist_ok=True)
    path = os.path.join(OUT_DIR, name)
    prs.save(path)
    print("saved:", path, os.path.getsize(path), "bytes,", len(prs.slides._sldIdLst), "slides")
    return path


# ================================================================ 方向 A：镜像
# 母题：两台设备左右对称。沿用前三套 deck 的模板语言（深色封面 + 浅底内容页
# + 白色卡片），好处是与既有三份汇报视觉一脉相承，像"同一个项目的第四份"。
# 差异点：封面右侧加入镜像双设备示意；内容页卡片用顶部色条而非左侧色条。

def build_a():
    prs = deck()

    # ---- A1 封面 ----
    s = blank(prs, NAVY)
    rect(s, 0, 0, 0.17, H, fill=BLUE)
    tbox(s, 1.05, 2.30, "ASSISTANT · 功能构想", 14, bold=True, color=BLUE)
    tbox(s, 1.05, 2.78, "设备同步", 44, bold=True, color=WHITE)
    tbox(s, 1.05, 3.98, "两台电脑共享文件 · 各自编辑 · 自动合并", 17, color=SOFT)
    rect(s, 1.05, 4.80, 3.0, 0.03, fill=RGBColor(0x35, 0x4A, 0x66))
    tbox(s, 1.05, 5.12, "软件工程导论实践 · 方案构想", 13, color=MUTED_L)

    # 右侧：两台设备镜像对称 + 中间同步符号
    by = 2.30
    for bx, label in [(9.05, "设备 A"), (11.15, "设备 B")]:
        rect(s, bx, by, 1.55, 1.15, fill=NAVY2, line_color=NAVY3, radius=0.08)
        rect(s, bx + 0.14, by + 0.17, 1.27, 0.62, fill=NAVY3)
        tbox(s, bx, by + 1.22, label, 11, bold=True, color=RGBColor(0x9F, 0xB2, 0xC8),
             align=PP_ALIGN.CENTER, maxw=1.55)
    tbox(s, 10.74, by + 0.30, "⇄", 26, bold=True, color=BLUE)
    # 框宽取两台设备的总跨度，让说明文字与设备整体居中对齐
    para(textbox(s, 9.05, by + 1.76, 3.65, 0.30), "同一份文件 · 两边都改 · 都不丢",
         size=12, color=SOFT2, first=True, align=PP_ALIGN.CENTER, after=0)
    pageno(s, 1, dark=True)

    # ---- A2 一页看懂 ----
    s = blank(prs)
    header(s, "OVERVIEW", "一页看懂")
    items = [
        ("要解决", "两台电脑的文件对不上", [
            "U 盘和微信传文件靠手工，改乱了不知道哪份新",
            "网盘冲突时静默生成副本，用户根本不知道发生过什么"], BLUE, BLUE_L, BLUE_TXT),
        ("怎么做", "把传输和同步逻辑拆开", [
            "同步逻辑自己写，传输做成可替换的通道",
            "局域网走直连，异地走共享目录，连不通自动降级"], TEAL, TEAL_L, TEAL_TXT),
        ("守什么", "冲突必须看得见", [
            "两边都改过的文件，保留双方版本，绝不静默覆盖",
            "删除的文件先移入回收目录，可恢复"], AMBER, AMBER_L, AMBER_TXT),
    ]
    for i, (tag, t, lines, ac, al, at) in enumerate(items):
        x = M + i * (CW / 3)
        w = CW / 3 - 0.25
        rect(s, x, BODY_TOP, w, 3.05, fill=WHITE, line_color=LINE, radius=0.05)
        rect(s, x, BODY_TOP, w, 0.06, fill=ac)
        chip(s, x + 0.32, BODY_TOP + 0.32, tag, al, at, 12)
        tbox(s, x + 0.32, BODY_TOP + 0.90, t, 16.5, bold=True, color=INK, maxw=w - 0.64)
        tf = textbox(s, x + 0.32, BODY_TOP + 1.55, w - 0.64, 1.3)
        for j, ln in enumerate(lines):
            para(tf, "·  " + ln, size=11.5, color=MUTED, first=(j == 0), after=8, line=1.3)
    pageno(s, 2)
    return save(prs, "draft_a_mirror.pptx")


# ================================================================ 方向 B：双栏
# 母题：左设备 / 中通道 / 右设备。把"两个对端 + 中间发生什么"直接做成版式骨架，
# 所有页面都按这个骨架组织，突出分布式结构本身。

def build_b():
    prs = deck()

    # ---- B1 封面：左右分栏，中间通道带 ----
    s = blank(prs, WHITE)
    GAP_L, GAP_R = 6.35, 7.00
    rect(s, 0, 0, GAP_L, H, fill=NAVY)              # 左：设备 A
    rect(s, GAP_R, 0, W - GAP_R, H, fill=NAVY4)     # 右：设备 B
    rect(s, GAP_L, 0, GAP_R - GAP_L, H, fill=BG)    # 中间通道带

    for i in range(7):
        c = BLUE if i % 2 == 0 else LINE
        rect(s, GAP_L + 0.07, 0.90 + i * 0.85, 0.51, 0.22, fill=c, radius=0.5)

    # 左栏文字：宽度一律限制在通道带左侧，避免压到浅色带上看不见
    tbox(s, 0.95, 2.15, "设备 A · 家里台式机", 12.5, bold=True, color=BLUE, maxw=5.0)
    tbox(s, 0.95, 3.95, "设备同步", 46, bold=True, color=WHITE, maxw=5.0)
    tbox(s, 0.95, 5.05, "两台电脑共享文件、各自编辑、自动合并", 16, color=SOFT, maxw=5.2)
    tbox(s, 0.95, 5.50, "冲突可见 · 版本不丢", 16, color=SOFT, maxw=5.2)
    tbox(s, 0.95, 6.05, "软件工程导论实践 · 方案构想", 12.5, color=MUTED_L, maxw=5.2)
    tbox(s, GAP_L + 0.055, 4.10, "⇄", 18, bold=True, color=BLUE, maxw=0.54)

    # 右栏：对端设备标注
    tbox(s, GAP_R + 0.55, 2.15, "设备 B · 公司笔记本", 12.5, bold=True,
         color=RGBColor(0x6F, 0xC5, 0xD6), maxw=5.2)

    # ---- B2 双栏架构页 ----
    s = blank(prs)
    header(s, "ARCHITECTURE", "一次同步，中间发生了什么")

    colw = 3.55
    left_x, mid_x, right_x = M, M + 4.10, M + 8.20

    def side(x, title, tag, accent, accent_l, accent_t, rows):
        rect(s, x, BODY_TOP, colw, 3.60, fill=WHITE, line_color=LINE, radius=0.05)
        rect(s, x, BODY_TOP, colw, 0.06, fill=accent)
        chip(s, x + 0.30, BODY_TOP + 0.28, tag, accent_l, accent_t, 11.5)
        tbox(s, x + 0.30, BODY_TOP + 0.80, title, 15.5, bold=True, color=INK,
             maxw=colw - 0.6)
        tf = textbox(s, x + 0.30, BODY_TOP + 1.42, colw - 0.6, 2.0)
        for j, (k, v) in enumerate(rows):
            para(tf, k, size=11.5, bold=True, color=accent_t, first=(j == 0),
                 after=1, line=1.2)
            para(tf, v, size=11, color=MUTED, after=9, line=1.25)

    side(left_x, "本地共享根", "本机", BLUE, BLUE_L, BLUE_TXT, [
        ("扫描 + 哈希", "算出每个文件的大小、时间、SHA-256"),
        ("比对基线", "和上次同步的记录比，判断谁改了"),
    ])
    side(right_x, "对端文件清单", "对端", TEAL, TEAL_L, TEAL_TXT, [
        ("拉取清单", "拿到对端每个文件的状态"),
        ("执行动作", "按判定结果上传 / 下载 / 标记冲突"),
    ])

    # 中间：三方差集判定
    rect(s, mid_x, BODY_TOP, colw, 3.60, fill=NAVY, radius=0.05)
    tbox(s, mid_x + 0.30, BODY_TOP + 0.28, "判定中心", 12, bold=True, color=BLUE,
         maxw=colw - 0.6)
    tbox(s, mid_x + 0.30, BODY_TOP + 0.72, "三方差集", 16, bold=True, color=WHITE,
         maxw=colw - 0.6)
    tf = textbox(s, mid_x + 0.30, BODY_TOP + 1.30, colw - 0.6, 0.80)
    para(tf, "本地 L  ·  基线 S  ·  对端 R", size=11.5, color=SOFT,
         first=True, after=6, line=1.3)
    para(tf, "三者一比，动作就唯一确定了", size=11, color=MUTED_L, after=0, line=1.3)

    for i, (cond, act, ac) in enumerate([("只有本地改", "上传", GREEN_CHIP),
                                         ("只有对端改", "下载", GREEN_CHIP),
                                         ("两边都改", "冲突", AMBER_CHIP)]):
        y = BODY_TOP + 2.22 + i * 0.44
        tbox(s, mid_x + 0.30, y + 0.06, cond, 11, color=SOFT, maxw=1.7)
        chip(s, mid_x + 2.05, y, act, ac, WHITE, 11)

    # 底部：通道降级
    rect(s, M, BODY_TOP + 3.85, CW, 0.95, fill=WHITE, line_color=LINE, radius=0.05)
    tbox(s, M + 0.32, BODY_TOP + 4.02, "通道选择（自动降级）", 12, bold=True,
         color=INK, maxw=3.2)
    _, _, w1, _ = chip(s, M + 0.32, BODY_TOP + 4.44, "直连 QTcpSocket", BLUE_L,
                       BLUE_TXT, 11)
    tbox(s, M + 0.32 + w1 + 0.16, BODY_TOP + 4.50, "▶", 12, color=MUTED)
    x2 = M + 0.32 + w1 + 0.16 + 0.34
    _, _, w2, _ = chip(s, x2, BODY_TOP + 4.44, "共享目录（异地）", TEAL_L,
                       TEAL_TXT, 11)
    tbox(s, x2 + w2 + 0.24, BODY_TOP + 4.52,
         "连不上直连就退回共享目录，所以两种网络环境都能用", 11, color=MUTED,
         maxw=6.4)
    pageno(s, 2)
    return save(prs, "draft_b_duo.pptx")


# ================================================================ 方向 C：状态流
# 母题：文件生命周期时间线。把"冲突"从缺陷变成流程上的一个正常节点，
# 让"不丢数据"这条底线成为画面的主角。
# 深色底上的状态文字统一用浅灰，颜色语义由色点承载 —— 彩色小字在
# 深蓝底上对比度只有 3.2~4.0，达不到 4.5 门槛。

def build_c():
    prs = deck()

    # ---- C1 封面：深底 + 贯穿时间线 ----
    s = blank(prs, NAVY)
    tbox(s, 1.05, 1.55, "ASSISTANT · 功能构想", 13, bold=True, color=BLUE)
    tbox(s, 1.05, 2.05, "一份文件的两台电脑之旅", 40, bold=True, color=WHITE)
    tbox(s, 1.05, 3.15, "设备同步 · 让每一处改动都留下痕迹，让每一次冲突都看得见",
         16, color=SOFT, maxw=10.0)

    ty = 4.55
    rect(s, 1.05, ty, 11.2, 0.03, fill=RGBColor(0x35, 0x4A, 0x66))
    stops = [("新建", GREEN, "已同步"), ("本地改", BLUE, "待上传"),
             ("对端也改", AMBER, "冲突"), ("保留两份", TEAL, "已解决")]
    for i, (act, ac, state) in enumerate(stops):
        cx = 1.05 + i * 3.73
        rect(s, cx, ty - 0.13, 0.20, 0.20, fill=ac, radius=0.5)
        w_act = tw(act, 12.5, True)
        tbox(s, cx + 0.10 - w_act / 2, ty - 0.72, act, 12.5, bold=True, color=WHITE)
        w_st = tw(state, 11)
        tbox(s, cx + 0.10 - w_st / 2, ty + 0.26, state, 11, color=SOFT)
    tbox(s, 1.05, 6.05, "软件工程导论实践 · 方案构想", 12.5, color=MUTED_L)
    pageno(s, 1, dark=True)

    # ---- C2 状态机页 ----
    s = blank(prs)
    header(s, "LIFECYCLE", "冲突不是故障，是流程上的一个站点")

    states = [
        ("与基线一致", "上次同步之后，谁都没动过", MUTED, ZEBRA),
        ("本地改了", "只在一侧改动 → 直接上传", GREEN, GREEN_L),
        ("对端改了", "只在另一侧改动 → 直接下载", GREEN, GREEN_L),
        ("两边都改了", "内容不同 → 生成冲突副本", AMBER, AMBER_L),
    ]
    bw, bh = 2.72, 1.55
    for i, (t, sub, ac, al) in enumerate(states):
        x = M + i * (bw + 0.31)
        rect(s, x, BODY_TOP, bw, bh, fill=WHITE, line_color=ac, line_w=1.25, radius=0.06)
        rect(s, x, BODY_TOP, bw, 0.07, fill=ac)
        tbox(s, x + 0.26, BODY_TOP + 0.32, t, 14.5, bold=True, color=INK,
             maxw=bw - 0.52)
        tf = textbox(s, x + 0.26, BODY_TOP + 0.82, bw - 0.52, 0.6)
        para(tf, sub, size=11, color=MUTED, first=True, after=0, line=1.3)
        if i < 3:
            tbox(s, x + bw + 0.03, BODY_TOP + 0.62, "▶", 12, color=MUTED)

    # 冲突处理：两份都留下
    rect(s, M, BODY_TOP + 2.00, CW, 1.42, fill=AMBER_L, line_color=AMBER,
         line_w=1.0, radius=0.05)
    tbox(s, M + 0.35, BODY_TOP + 2.18, "冲突发生时：两份都留下，一份都不丢",
         15, bold=True, color=AMBER_TXT)
    tf = textbox(s, M + 0.35, BODY_TOP + 2.62, CW - 0.7, 0.70)
    para(tf, "原文件保持不动，对端版本另存为「文件名 (冲突副本 来自设备名 时间戳).扩展名」，"
             "并在界面里列出来 —— 由用户自己决定留哪份，系统不做仲裁。",
         size=11.5, color=RGBColor(0x6B, 0x42, 0x00), first=True, after=0, line=1.35)

    # 底部两条底线
    for i, (t, sub) in enumerate([
            ("删除也可恢复", "对端删除的文件先移入回收目录，不是直接消失"),
            ("同步可被打断", "中途中断后再同步能正确收敛，不留重复和残留")]):
        x = M + i * (CW / 2 + 0.12)
        w = CW / 2 - 0.12
        rect(s, x, BODY_TOP + 3.62, w, 0.92, fill=WHITE, line_color=LINE, radius=0.05)
        rect(s, x, BODY_TOP + 3.62, 0.055, 0.92, fill=TEAL)
        tbox(s, x + 0.32, BODY_TOP + 3.80, t, 13, bold=True, color=INK, maxw=w - 0.6)
        tbox(s, x + 0.32, BODY_TOP + 4.14, sub, 11, color=MUTED, maxw=w - 0.6)
    pageno(s, 2)
    return save(prs, "draft_c_flow.pptx")


if __name__ == "__main__":
    build_a()
    build_b()
    build_c()
