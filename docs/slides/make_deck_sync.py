# -*- coding: utf-8 -*-
"""Assistant 设备同步模块 —— 【开题汇报】PPT 生成脚本（11 页，构想为主、技术细节从简）

用法： python make_deck_sync.py
产出： docs/slides/设备同步模块-开题汇报.pptx（同目录）

视觉方向：A「镜像」—— 沿用文件检索 / 文件管理两套 deck 的模板语言
（深蓝封面 NAVY #142238 + 浅底内容页 + 白色卡片 + 胶囊标签），
封面与结尾用左右对称的双设备表达"两端"。

两条硬约束（由 check_deck_geometry.py / check_deck_contrast.py 验证）：
  1. 文本框宽度一律用 tw() 估算后设定，不"猜"——框宽过大会让文字压到
     不该压的区域，也会让对比度自检因中心点落在错误背景上而误报。
  2. 文字颜色用 *_TXT 深色变体，色块填充用 *_CHIP 加深变体。
     主色（BLUE/GREEN/AMBER/TEAL）是为图形设计的，直接当小字或当
     白字底色都达不到 4.5:1。
"""
import os

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

FONT = "Microsoft YaHei"

# ---- 图形/色块用色（沿用既有模板）----
NAVY    = RGBColor(0x14, 0x22, 0x38)
NAVY2   = RGBColor(0x1E, 0x30, 0x4C)
NAVY3   = RGBColor(0x2A, 0x3F, 0x5E)
DARKLN  = RGBColor(0x35, 0x4A, 0x66)
BLUE    = RGBColor(0x4A, 0x90, 0xD9)
BLUE_L  = RGBColor(0xE8, 0xF1, 0xFB)
TEAL    = RGBColor(0x1F, 0x7A, 0x8C)
TEAL_L  = RGBColor(0xE4, 0xF2, 0xF4)
GREEN   = RGBColor(0x2E, 0x8B, 0x57)
GREEN_L = RGBColor(0xE7, 0xF4, 0xEC)
AMBER   = RGBColor(0xB5, 0x6E, 0x00)
AMBER_L = RGBColor(0xFD, 0xF3, 0xE2)
RED     = RGBColor(0xC0, 0x39, 0x2B)
RED_L   = RGBColor(0xFB, 0xEC, 0xEA)

INK     = RGBColor(0x1A, 0x1A, 0x1A)
MUTED   = RGBColor(0x5E, 0x6A, 0x7A)
MUTED_L = RGBColor(0x8E, 0x9A, 0xAD)
LINE    = RGBColor(0xD8, 0xDE, 0xE6)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
BG      = RGBColor(0xF7, 0xF9, 0xFC)
ZEBRA   = RGBColor(0xF1, 0xF5, 0xFA)
SOFT    = RGBColor(0xB9, 0xC6, 0xD8)
SOFT2   = RGBColor(0xC7, 0xD4, 0xE6)

# ---- 文字专用深色变体（>=4.5:1）----
BLUE_TXT  = RGBColor(0x2C, 0x5F, 0x9E)
TEAL_TXT  = RGBColor(0x17, 0x62, 0x70)
GREEN_TXT = RGBColor(0x26, 0x72, 0x4A)
AMBER_TXT = RGBColor(0x8A, 0x54, 0x00)
RED_TXT   = RGBColor(0xA8, 0x2E, 0x22)

# ---- 色块填充加深变体（白字压主色只有 3.3~4.3:1）----
#      TEAL 填充实测 4.98:1，本来就够，不动。
GREEN_CHIP = RGBColor(0x2A, 0x7D, 0x50)
AMBER_CHIP = RGBColor(0x9C, 0x5D, 0x00)
BLUE_CHIP  = RGBColor(0x2C, 0x5F, 0x9E)

W, H = 13.333, 7.5
M, CW = 0.75, 11.833
BODY_TOP = 1.80

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "设备同步模块-开题汇报.pptx")


# ---------------------------------------------------------------- 尺寸估算

def tw(text, size, bold=False):
    """估算一段文字单行渲染所需宽度（英寸）。中文 1.0 em / 英文数字 0.55 em。"""
    best = 0.0
    for line in str(text).split("\n"):
        cjk = sum(1 for ch in line if ord(ch) > 0x2E80)
        other = len(line) - cjk
        w = (cjk * size + other * size * 0.55) * (1.03 if bold else 1.0)
        best = max(best, w)
    return best / 72.0 + 0.08


# ---------------------------------------------------------------- 版式基础

def style_font(run, size=None, bold=None, color=None, name=FONT, italic=None):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ('a:ea', 'a:cs'):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set('typeface', name)
    if size   is not None: run.font.size  = Pt(size)
    if bold   is not None: run.font.bold  = bold
    if italic is not None: run.font.italic = italic
    if color  is not None: run.font.color.rgb = color


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def para(tf, text, size=14, bold=False, color=INK, first=False, before=0, after=5,
         align=PP_ALIGN.LEFT, line=1.2, italic=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(before); p.space_after = Pt(after); p.line_spacing = line
    r = p.add_run(); r.text = text
    style_font(r, size=size, bold=bold, color=color, italic=italic)
    return p


def tbox(slide, x, y, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT,
         maxw=None, h=None):
    """宽度按文字实际长度设定的文本框（首选，避免框宽远大于文字）。"""
    w = tw(text, size, bold)
    if maxw is not None:
        w = min(w, maxw)
    if h is None:
        h = size * 1.45 / 72.0
    tf = textbox(slide, x, y, w, h)
    para(tf, text, size=size, bold=bold, color=color, first=True,
         align=align, after=0)
    return tf


def rect(slide, x, y, w, h, fill=None, line_color=None, line_w=0.75, radius=None):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None: shp.fill.background()
    else:            shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line_color is None: shp.line.fill.background()
    else:                  shp.line.color.rgb = line_color; shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    if radius: shp.adjustments[0] = radius
    return shp


def chip(slide, x, y, text, fill, txt_color, size=12, pad=0.30, h=0.36, radius=0.4):
    """按文字长度自适应的胶囊标签。"""
    w = tw(text, size, True) + pad * 2
    rect(slide, x, y, w, h, fill=fill, radius=radius)
    para(textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.MIDDLE), text, size=size,
         bold=True, color=txt_color, first=True, align=PP_ALIGN.CENTER, after=0)
    return x, y, w, h


def blank(prs, bg=BG):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, fill=bg)
    return s


def header(slide, kicker, title):
    tbox(slide, M, 0.42, kicker, 12, bold=True, color=BLUE_TXT)
    tbox(slide, M, 0.76, title, 27, bold=True, color=INK)
    rect(slide, M, 1.50, 1.05, 0.045, fill=BLUE)


def pageno(slide, n, dark=False):
    t = f"{n:02d}"
    tbox(slide, W - M - tw(t, 11, True), 6.98, t, 11, bold=True,
         color=MUTED_L if dark else MUTED)


def table(slide, x, y, rows, col_w, row_h=0.72, head_h=0.44, fs=12.5, head_fs=12):
    n_r, n_c = len(rows), len(rows[0])
    total_w = sum(col_w)
    gt = slide.shapes.add_table(n_r, n_c, Inches(x), Inches(y), Inches(total_w),
                                Inches(head_h + row_h * (n_r - 1))).table
    gt.first_row = False; gt.horz_banding = False
    for i, cw in enumerate(col_w):
        gt.columns[i].width = Inches(cw)
    gt.rows[0].height = Inches(head_h)
    for r in range(1, n_r):
        gt.rows[r].height = Inches(row_h)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = gt.cell(r, c)
            cell.margin_left = Inches(0.15); cell.margin_right = Inches(0.13)
            cell.margin_top = Inches(0.04);  cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if r == 0 else (WHITE if r % 2 else ZEBRA)
            tf = cell.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; p.line_spacing = 1.15
            run = p.add_run(); run.text = str(val)
            style_font(run, size=head_fs if r == 0 else fs,
                       bold=(r == 0), color=WHITE if r == 0 else INK)
    return gt


def bullets(slide, x, y, w, h, items, size=11.5, color=MUTED, gap=8, line=1.32,
            marker="·  "):
    tf = textbox(slide, x, y, w, h)
    for i, it in enumerate(items):
        para(tf, marker + it, size=size, color=color, first=(i == 0),
             after=gap, line=line)
    return tf


# ================================================================ 页面

def p01_cover(prs):
    s = blank(prs, NAVY)
    rect(s, 0, 0, 0.17, H, fill=BLUE)
    tbox(s, 1.05, 2.30, "ASSISTANT · 功能构想", 14, bold=True, color=BLUE)
    tbox(s, 1.05, 2.78, "设备同步", 44, bold=True, color=WHITE)
    tbox(s, 1.05, 3.98, "两台电脑共享文件 · 各自编辑 · 自动合并", 17, color=SOFT)
    rect(s, 1.05, 4.80, 3.0, 0.03, fill=DARKLN)
    tbox(s, 1.05, 5.12, "软件工程导论实践 · 开题汇报", 13, color=MUTED_L)

    by = 2.30
    for bx, label in [(9.05, "设备 A"), (11.15, "设备 B")]:
        rect(s, bx, by, 1.55, 1.15, fill=NAVY2, line_color=NAVY3, radius=0.08)
        rect(s, bx + 0.14, by + 0.17, 1.27, 0.62, fill=NAVY3)
        tbox(s, bx, by + 1.22, label, 11, bold=True,
             color=RGBColor(0x9F, 0xB2, 0xC8), align=PP_ALIGN.CENTER, maxw=1.55)
    tbox(s, 10.74, by + 0.30, "⇄", 26, bold=True, color=BLUE)
    para(textbox(s, 9.05, by + 1.76, 3.65, 0.30), "同一份文件 · 两边都改 · 都不丢",
         size=12, color=SOFT2, first=True, align=PP_ALIGN.CENTER, after=0)
    pageno(s, 1, dark=True)


def p02_overview(prs):
    s = blank(prs)
    header(s, "OVERVIEW", "一页看懂")
    items = [
        ("要解决", "两台电脑的文件对不上", [
            "U 盘和微信传文件靠手工，改乱了不知道哪份新",
            "网盘冲突时静默生成副本，用户根本不知道发生过什么"],
         BLUE, BLUE_L, BLUE_TXT),
        ("怎么做", "把传输和同步逻辑拆开", [
            "同步逻辑自己写，传输做成可替换的通道",
            "局域网走直连，异地走共享目录，连不通自动降级"],
         TEAL, TEAL_L, TEAL_TXT),
        ("守什么", "冲突必须看得见", [
            "两边都改过的文件，保留双方版本，绝不静默覆盖",
            "删除的文件先移入回收目录，可恢复"],
         AMBER, AMBER_L, AMBER_TXT),
    ]
    for i, (tag, t, lines, ac, al, at) in enumerate(items):
        x = M + i * (CW / 3)
        w = CW / 3 - 0.25
        rect(s, x, BODY_TOP, w, 3.05, fill=WHITE, line_color=LINE, radius=0.05)
        rect(s, x, BODY_TOP, w, 0.06, fill=ac)
        chip(s, x + 0.32, BODY_TOP + 0.32, tag, al, at, 12)
        tbox(s, x + 0.32, BODY_TOP + 0.90, t, 16.5, bold=True, color=INK,
             maxw=w - 0.64)
        bullets(s, x + 0.32, BODY_TOP + 1.55, w - 0.64, 1.3, lines)
    pageno(s, 2)


def p03_problem(prs):
    s = blank(prs)
    header(s, "WHY", "现在是怎么做的，为什么不够")
    rows = [
        ["现有做法", "它解决了什么", "它哪里不够"],
        ["U 盘 / 微信传文件", "能把文件挪到另一台机器",
         "全靠手工；两边都改过之后，分不清哪份最新"],
        ["网盘同步盘", "自动同步，不用管",
         "冲突时静默生成副本，用户不知道发生过什么；也无法按需选择共享范围"],
        ["Syncthing 之类的工具", "点对点同步，功能很强",
         "面向技术人员，配置复杂，与电脑里已有的工具是割裂的"],
    ]
    table(s, M, BODY_TOP, rows, [3.05, 3.55, 5.233], row_h=0.78)

    yy = BODY_TOP + 0.44 + 3 * 0.78 + 0.42
    rect(s, M, yy, CW, 0.88, fill=AMBER_L, line_color=AMBER, line_w=1.0, radius=0.05)
    tbox(s, M + 0.35, yy + 0.17, "三种做法的共同盲点", 13.5, bold=True,
         color=AMBER_TXT)
    tbox(s, M + 0.35, yy + 0.48,
         "都没有把「冲突」当成一件需要让用户看见的事 —— 要么静默覆盖，要么悄悄生成一个副本。",
         11.5, color=RGBColor(0x6B, 0x42, 0x00), maxw=CW - 0.7)
    pageno(s, 3)


def p04_goal(prs):
    s = blank(prs)
    header(s, "GOAL", "我们要做到什么")

    rect(s, M, BODY_TOP, CW, 0.82, fill=NAVY, radius=0.05)
    rect(s, M, BODY_TOP, 0.055, 0.82, fill=BLUE)
    tbox(s, M + 0.38, BODY_TOP + 0.24,
         "让两台电脑上的同一批文件，最终一致；两边都改过时，两个版本都留下",
         15.5, bold=True, color=WHITE, maxw=CW - 0.76)

    y2 = BODY_TOP + 1.12
    colw = CW / 2 - 0.16
    for i, (tag, ac, al, at, items) in enumerate([
        ("功能上", BLUE, BLUE_L, BLUE_TXT, [
            "两台设备配对后可互相同步指定目录",
            "新增、修改的文件在两端保持一致",
            "两端都改过时生成冲突副本，绝不静默覆盖",
            "删除的文件先移入回收目录，可恢复",
            "支持手动同步与定时自动同步"]),
        ("质量上", TEAL, TEAL_L, TEAL_TXT, [
            "局域网与异地两种网络环境都能用",
            "冲突双方版本都必须留存，一个都不能丢",
            "同步在后台进行，不卡界面",
            "单台电脑开两个实例就能完整演示"]),
    ]):
        x = M + i * (CW / 2 + 0.32)
        rect(s, x, y2, colw, 2.80, fill=WHITE, line_color=LINE, radius=0.05)
        rect(s, x, y2, colw, 0.06, fill=ac)
        chip(s, x + 0.30, y2 + 0.28, tag, al, at, 12)
        bullets(s, x + 0.30, y2 + 0.88, colw - 0.6, 1.90, items)
    pageno(s, 4)


def p05_decision(prs):
    s = blank(prs)
    header(s, "KEY DECISION", "把「传输」和「同步逻辑」拆开")

    # 左：架构示意
    lw = 5.55
    rect(s, M, BODY_TOP, lw, 0.86, fill=NAVY, radius=0.05)
    tbox(s, M + 0.34, BODY_TOP + 0.26, "同步引擎（自己写）", 15, bold=True,
         color=WHITE)
    tbox(s, M + 0.34, BODY_TOP + 0.56, "比对、判定、冲突处理", 10.5, color=SOFT)

    for i, (t, sub, ac, al, at) in enumerate([
            ("直连通道", "同一局域网时用，快", BLUE, BLUE_L, BLUE_TXT),
            ("共享目录通道", "异地时用，靠网盘搬运", TEAL, TEAL_L, TEAL_TXT)]):
        x = M + i * (lw / 2 + 0.16)
        w = lw / 2 - 0.16
        yy = BODY_TOP + 1.42
        rect(s, x, yy, w, 0.92, fill=WHITE, line_color=ac, line_w=1.25, radius=0.05)
        tbox(s, x + 0.26, yy + 0.16, t, 13.5, bold=True, color=at, maxw=w - 0.5)
        tbox(s, x + 0.26, yy + 0.50, sub, 10.5, color=MUTED, maxw=w - 0.5)

    tbox(s, M + lw / 2 - 0.11, BODY_TOP + 1.02, "▼", 13, color=MUTED)

    # 关键点：实现可替换
    rect(s, M, BODY_TOP + 2.58, lw, 1.20, fill=BLUE_L, line_color=BLUE,
         line_w=1.0, radius=0.05)
    tbox(s, M + 0.32, BODY_TOP + 2.74, "两个通道实现同一个接口", 13, bold=True,
         color=BLUE_TXT)
    tbox(s, M + 0.32, BODY_TOP + 3.06,
         "换通道不影响同步逻辑一行代码；跑不通就自动退回另一个",
         11, color=RGBColor(0x24, 0x4A, 0x7A), maxw=lw - 0.64)

    # 右：为什么
    rx = M + lw + 0.42
    rw = CW - lw - 0.42
    tbox(s, rx, BODY_TOP + 0.04, "为什么这样拆：两种网络环境的代价差太远",
         13.5, bold=True, color=INK, maxw=rw)
    for i, (t, sub, ac, al, at) in enumerate([
            ("同一局域网", "自己写直连是可控的工作量", GREEN, GREEN_L, GREEN_TXT),
            ("异地（跨网络）", "要额外解决端口映射、中继服务器，成本和运维都不可控",
             RED, RED_L, RED_TXT)]):
        yy = BODY_TOP + 0.50 + i * 1.32
        rect(s, rx, yy, rw, 1.16, fill=WHITE, line_color=LINE, radius=0.05)
        rect(s, rx, yy, 0.055, 1.16, fill=ac)
        tbox(s, rx + 0.30, yy + 0.16, t, 13, bold=True, color=at, maxw=rw - 0.6)
        bullets(s, rx + 0.30, yy + 0.50, rw - 0.6, 0.6, [sub], size=11, gap=0)

    yy = BODY_TOP + 3.14
    rect(s, rx, yy, rw, 0.64, fill=ZEBRA, line_color=LINE, radius=0.05)
    tbox(s, rx + 0.30, yy + 0.19, "结论：异地交给网盘，自己只写同步逻辑",
         12, bold=True, color=INK, maxw=rw - 0.6)
    pageno(s, 5)


def p06_flow(prs):
    s = blank(prs)
    header(s, "FLOW", "一次同步，中间发生了什么")

    steps = [
        ("1", "扫描本地", ["算出每个文件的", "大小、时间、哈希"], BLUE_CHIP, BLUE_TXT),
        ("2", "拉取清单", ["拿到对端每个", "文件的状态"], TEAL, TEAL_TXT),
        ("3", "比对基线", ["和上次同步的", "记录做比较"], BLUE_CHIP, BLUE_TXT),
        ("4", "执行动作", ["上传 / 下载 /", "标记为冲突"], AMBER_CHIP, AMBER_TXT),
    ]
    gap = 0.34
    bw = (CW - gap * 3) / 4
    for i, (num, t, lines, ac, at) in enumerate(steps):
        x = M + i * (bw + gap)
        rect(s, x, BODY_TOP, bw, 2.30, fill=WHITE, line_color=LINE, radius=0.05)
        rect(s, x, BODY_TOP, bw, 0.06, fill=ac)
        chip(s, x + 0.30, BODY_TOP + 0.30, num, ac, WHITE, 12, pad=0.24, h=0.38)
        tbox(s, x + 0.30, BODY_TOP + 0.86, t, 15.5, bold=True, color=INK,
             maxw=bw - 0.6)
        bullets(s, x + 0.30, BODY_TOP + 1.40, bw - 0.6, 0.8, lines, size=11, gap=3)
        if i < 3:
            para(textbox(s, x + bw, BODY_TOP + 0.98, gap, 0.36,
                         anchor=MSO_ANCHOR.MIDDLE), "▶", size=13, color=MUTED,
                 first=True, align=PP_ALIGN.CENTER, after=0)

    yy = BODY_TOP + 2.62
    rect(s, M, yy, CW, 0.92, fill=NAVY, radius=0.05)
    tbox(s, M + 0.36, yy + 0.16, "关键在第三步", 13, bold=True, color=BLUE)
    tbox(s, M + 0.36, yy + 0.47,
         "判定只看「谁改了」——只有一边改就直接同步；两边都改过才需要人来决定。",
         12, color=SOFT2, maxw=CW - 0.72)
    pageno(s, 6)


def p07_conflict(prs):
    s = blank(prs)
    header(s, "CONFLICT", "两边都改了，怎么办")

    # 上：两个版本
    cw2 = 5.2
    for i, (who, note, ac, al, at, name) in enumerate([
            ("你改的版本", "在家里的台式机上改的", BLUE, BLUE_L, BLUE_TXT,
             "季度报告.docx"),
            ("同事改的版本", "在公司笔记本上改的", TEAL, TEAL_L, TEAL_TXT,
             "季度报告.docx")]):
        x = M + i * (CW - cw2)
        rect(s, x, BODY_TOP, cw2, 1.42, fill=WHITE, line_color=ac, line_w=1.25,
             radius=0.05)
        rect(s, x, BODY_TOP, cw2, 0.06, fill=ac)
        tbox(s, x + 0.30, BODY_TOP + 0.26, who, 14, bold=True, color=at)
        tbox(s, x + 0.30, BODY_TOP + 0.60, name, 12, color=INK)
        tbox(s, x + 0.30, BODY_TOP + 0.92, note, 11, color=MUTED, maxw=cw2 - 0.6)

    # 汇聚箭头
    cx = W / 2
    rect(s, M + cw2 / 2 + 0.1, BODY_TOP + 1.42, cx - 0.1 - (M + cw2 / 2), 0.02,
         fill=LINE)
    rect(s, cx + 0.1, BODY_TOP + 1.42, W - M - cw2 / 2 - 0.1 - cx, 0.02, fill=LINE)
    tbox(s, cx - 0.09, BODY_TOP + 1.56, "▼", 13, color=MUTED)

    # 下：结果
    yy = BODY_TOP + 2.02
    rect(s, 2.4, yy, CW - 3.3, 1.72, fill=AMBER_L, line_color=AMBER, line_w=1.25,
         radius=0.05)
    tbox(s, 2.72, yy + 0.18, "两份都保留，一份都不丢", 16, bold=True, color=AMBER_TXT)
    # label 列与 value 列固定起点，避免因标签字数不同而错位
    tbox(s, 2.72, yy + 0.62, "原文件不动", 12, color=RGBColor(0x6B, 0x42, 0x00))
    tbox(s, 4.05, yy + 0.62, "季度报告.docx", 12, bold=True,
         color=RGBColor(0x6B, 0x42, 0x00))
    tbox(s, 2.72, yy + 0.98, "对端版本另存", 12, color=RGBColor(0x6B, 0x42, 0x00))
    tbox(s, 4.05, yy + 0.98,
         "季度报告 (冲突副本 来自笔记本 20260924-1530).docx", 12, bold=True,
         color=RGBColor(0x6B, 0x42, 0x00))

    tbox(s, M, BODY_TOP + 4.00, "由用户决定留哪份，系统不做仲裁", 14, bold=True,
         color=INK)
    tbox(s, M, BODY_TOP + 4.38,
         "冲突会出现在界面的冲突列表里，可以一键打开所在文件夹对照。",
         11.5, color=MUTED, maxw=CW)
    pageno(s, 7)


def p08_baseline(prs):
    s = blank(prs)
    header(s, "BASELINE", "两条不能让的底线")

    colw = CW / 2 - 0.16
    for i, (tag, t, ac, al, at, rows) in enumerate([
        ("不丢数据", "任何情况下都不静默覆盖", AMBER, AMBER_L, AMBER_TXT, [
            ("两边都改过", "保留双方版本，生成冲突副本"),
            ("对端删除了文件", "先移入回收目录，不是直接消失"),
            ("同步被中断", "再同步能正确收敛，不留重复和残留")]),
        ("不给坏人开门", "接收端不信任对端送来的路径", BLUE, BLUE_L, BLUE_TXT, [
            ("路径校验", "拒绝一切越出共享目录的路径，包括符号链接绕道"),
            ("配对确认", "首次配对要两台机器上人工点确认"),
            ("只收该收的", "共享目录之外的位置一概不写")]),
    ]):
        x = M + i * (CW / 2 + 0.32)
        rect(s, x, BODY_TOP, colw, 3.60, fill=WHITE, line_color=LINE, radius=0.05)
        rect(s, x, BODY_TOP, colw, 0.06, fill=ac)
        chip(s, x + 0.30, BODY_TOP + 0.28, tag, al, at, 12)
        tbox(s, x + 0.30, BODY_TOP + 0.80, t, 15.5, bold=True, color=INK,
             maxw=colw - 0.6)
        for j, (k, v) in enumerate(rows):
            yy = BODY_TOP + 1.42 + j * 0.72
            tbox(s, x + 0.30, yy, k, 11.5, bold=True, color=at, maxw=colw - 0.6)
            tbox(s, x + 0.30, yy + 0.28, v, 11, color=MUTED, maxw=colw - 0.6)
    pageno(s, 8)


def p09_feasibility(prs):
    s = blank(prs)
    header(s, "FEASIBILITY", "技术上可行，而且能当场演示")

    lw = 7.15
    rect(s, M, BODY_TOP, lw, 3.55, fill=WHITE, line_color=LINE, radius=0.05)
    rect(s, M, BODY_TOP, lw, 0.06, fill=BLUE)
    tbox(s, M + 0.32, BODY_TOP + 0.30, "需要的能力，项目里全都有", 15.5,
         bold=True, color=INK)

    caps = [
        ("网络通信（TCP / HTTP）", "Qt6::Network", "已链接"),
        ("文件哈希校验", "QCryptographicHash", "Qt 内置"),
        ("同步元数据存储", "SQLite", "已有 4 个库"),
        ("设置与状态界面", "QML 主面板", "现成可复用"),
        ("后台同步线程", "QThread", "Qt 内置"),
    ]
    for i, (cap, lib, st) in enumerate(caps):
        yy = BODY_TOP + 0.92 + i * 0.50
        if i % 2 == 0:
            rect(s, M + 0.20, yy - 0.05, lw - 0.4, 0.44, fill=ZEBRA, radius=0.3)
        tbox(s, M + 0.38, yy + 0.05, cap, 12, color=INK, maxw=3.6)
        tbox(s, M + 4.15, yy + 0.05, lib, 11.5, color=MUTED, maxw=2.2)
        chip(s, M + lw - 1.55, yy - 0.01, st, GREEN_L, GREEN_TXT, 10.5,
             pad=0.22, h=0.32)

    rx = M + lw + 0.42
    rw = CW - lw - 0.42
    rect(s, rx, BODY_TOP, rw, 3.55, fill=NAVY, radius=0.05)
    tbox(s, rx + 0.30, BODY_TOP + 0.30, "演示方案", 13, bold=True, color=BLUE)
    tbox(s, rx + 0.30, BODY_TOP + 0.70, "一台电脑就够", 17, bold=True, color=WHITE,
         maxw=rw - 0.6)
    bullets(s, rx + 0.30, BODY_TOP + 1.32, rw - 0.6, 1.4, [
        "开两个实例，各用一个配置目录",
        "两个实例指向同一个共享目录",
        "现场制造一次冲突，看两份文件都在"], size=11, color=SOFT2, gap=8)
    rect(s, rx + 0.30, BODY_TOP + 2.90, rw - 0.6, 0.02, fill=DARKLN)
    tbox(s, rx + 0.30, BODY_TOP + 3.04, "不需要第二台电脑或可用网络", 11.5,
         color=SOFT)
    pageno(s, 9)


def p10_plan(prs):
    s = blank(prs)
    header(s, "PLAN", "四个人，六个里程碑")

    gap = 0.25
    bw = (CW - gap * 3) / 4
    people = [
        ("同学 A", "网络与传输", ["直连通道与协议", "断点续传"], BLUE, BLUE_TXT),
        ("同学 B", "同步引擎", ["差量判定算法", "冲突策略"], TEAL, TEAL_TXT),
        ("同学 C", "数据与界面", ["数据库与桥接", "主面板界面"], BLUE, BLUE_TXT),
        ("同学 D", "安全与测试", ["路径校验", "集成测试与文档"], AMBER, AMBER_TXT),
    ]
    for i, (who, t, lines, ac, at) in enumerate(people):
        x = M + i * (bw + gap)
        rect(s, x, BODY_TOP, bw, 1.92, fill=WHITE, line_color=LINE, radius=0.05)
        rect(s, x, BODY_TOP, bw, 0.06, fill=ac)
        tbox(s, x + 0.26, BODY_TOP + 0.28, who, 11.5, bold=True, color=at)
        tbox(s, x + 0.26, BODY_TOP + 0.60, t, 14, bold=True, color=INK,
             maxw=bw - 0.52)
        bullets(s, x + 0.26, BODY_TOP + 1.10, bw - 0.52, 0.7, lines, size=10.5,
                gap=4)

    # 里程碑时间线：首尾各留 1 英寸，否则两端节点的文字框会顶到内容边界
    ty = BODY_TOP + 2.62
    x0, x1 = M + 1.0, W - M - 1.0
    rect(s, x0 - 0.5, ty, (x1 - x0) + 1.0, 0.02, fill=LINE)
    miles = [("1", "需求与\n契约冻结"), ("2", "四大模块\n并行开发"),
             ("3", "引擎\n集成"), ("4", "端到端\n与安全"),
             ("5", "完善与\n打磨"), ("6", "文档与\n答辩准备")]
    span = (x1 - x0) / (len(miles) - 1)
    for i, (n, label) in enumerate(miles):
        cx = x0 + i * span
        rect(s, cx - 0.135, ty - 0.12, 0.27, 0.27, fill=BLUE_CHIP, radius=0.5)
        para(textbox(s, cx - 0.135, ty - 0.11, 0.27, 0.27, anchor=MSO_ANCHOR.MIDDLE),
             n, size=10, bold=True, color=WHITE, first=True,
             align=PP_ALIGN.CENTER, after=0)
        tf = textbox(s, cx - 0.75, ty + 0.26, 1.5, 0.6)
        for j, ln in enumerate(label.split("\n")):
            para(tf, ln, size=10.5, color=MUTED, first=(j == 0), after=1,
                 align=PP_ALIGN.CENTER, line=1.2)

    yy = ty + 1.12
    rect(s, M, yy, CW, 0.66, fill=BLUE_L, line_color=BLUE, line_w=1.0, radius=0.05)
    tbox(s, M + 0.32, yy + 0.19,
         "前两周全员一起冻结接口，之后四个人才能真正并行 —— 这是整个计划里最关键的一步。",
         12, bold=True, color=BLUE_TXT, maxw=CW - 0.64)
    pageno(s, 10)


def p11_end(prs):
    s = blank(prs, NAVY)
    rect(s, 0, 0, 0.17, H, fill=BLUE)
    tbox(s, 1.05, 2.05, "我们想做的", 14, bold=True, color=BLUE)
    tbox(s, 1.05, 2.52, "一个能看见冲突、并且不丢数据的同步", 34, bold=True,
         color=WHITE, maxw=11.2)
    tbox(s, 1.05, 3.62,
         "不是把文件悄悄挪过去，而是让每一处改动都留下痕迹，让用户始终知道发生了什么。",
         15, color=SOFT, maxw=11.0)
    rect(s, 1.05, 4.35, 3.0, 0.03, fill=DARKLN)

    tbox(s, 1.05, 4.70, "请老师指点", 13, bold=True, color=BLUE)
    bullets(s, 1.05, 5.10, 11.2, 1.5, [
        "共享目录只做一个，还是允许配置多个",
        "自动同步的默认间隔取多少合适",
        "要不要把实时协同编辑也纳入范围"], size=12, color=SOFT2, gap=8)
    pageno(s, 11, dark=True)


# ================================================================ main

if __name__ == "__main__":
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    for fn in (p01_cover, p02_overview, p03_problem, p04_goal, p05_decision,
               p06_flow, p07_conflict, p08_baseline, p09_feasibility,
               p10_plan, p11_end):
        fn(prs)
    prs.save(OUT)
    print("saved:", OUT, os.path.getsize(OUT), "bytes,",
          len(prs.slides._sldIdLst), "slides")
