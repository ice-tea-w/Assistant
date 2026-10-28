# -*- coding: utf-8 -*-
"""一次性脚本：生成演示用合同（docx）到工程根目录（随工程拷贝即可换机使用）。"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "合同.docx"

doc = Document()

style = doc.styles["Normal"]
style.font.name = "宋体"
style.font.size = Pt(11)
style.element.rPr.rFonts.set(
    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "宋体"
)

title = doc.add_heading("技术开发（委托）合同", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph("合同编号：HT-2026-1022")
doc.add_paragraph("甲方（委托方）：星河科技有限公司")
doc.add_paragraph("乙方（受托方）：远山软件工作室")
doc.add_paragraph("")

doc.add_heading("第一条 项目内容", level=1)
doc.add_paragraph(
    "甲方委托乙方开发桌面助手文件安全模块，包括文件加密、敏感文件检测、"
    "操作审计与安全备份四项功能。乙方应按双方确认的需求文档与排期交付。"
)

doc.add_heading("第二条 合同金额与付款方式", level=1)
doc.add_paragraph("合同总金额为人民币 40,000 元（大写：肆万元整）。")
doc.add_paragraph("付款方式：合同签订后 5 个工作日内支付 50%，验收通过后支付剩余 50%。")

doc.add_heading("第三条 保密条款", level=1)
doc.add_paragraph(
    "双方对合作过程中获知的商业秘密、技术资料、用户数据承担保密义务，"
    "未经对方书面同意不得向任何第三方披露。本条款在合同终止后两年内继续有效。"
)
doc.add_paragraph("内部资料，请勿外传。")

doc.add_heading("第四条 交付与验收", level=1)
doc.add_paragraph(
    "乙方应于 2026 年 12 月 1 日前完成交付，甲方在 10 个工作日内组织验收；"
    "验收标准以需求文档中的功能清单与安全验收清单为准。"
)

doc.add_heading("第五条 联系方式", level=1)
doc.add_paragraph("联系人：张三    电话：13800138000    邮箱：zhangsan@example.com")

doc.add_paragraph("")
doc.add_paragraph("甲方（盖章）：____________________        乙方（盖章）：____________________")
doc.add_paragraph("签署日期：2026 年 10 月 22 日          签署日期：2026 年 10 月 22 日")

doc.save(OUTPUT)
print(f"saved {OUTPUT}")
